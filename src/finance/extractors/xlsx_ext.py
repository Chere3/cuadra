# -*- coding: utf-8 -*-
"""xlsx_ext.py — Extractor de estados en XLSX. Hoja 'meta' (key/value) opcional + hoja de datos."""
from __future__ import annotations
import openpyxl
from . import EXTRACTOR_VERSION
from .csv_ext import _canon

META_KEYS = {"account_id", "currency", "period_start", "period_end",
             "opening_balance", "closing_balance", "declared_credits",
             "declared_debits", "declared_count"}


def extract_xlsx(path, account_hint=None, policy=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    meta = {"declared": {}}
    # hoja meta
    if "meta" in [s.lower() for s in wb.sheetnames]:
        ws = wb[[s for s in wb.sheetnames if s.lower() == "meta"][0]]
        for r in ws.iter_rows(values_only=True):
            if not r or r[0] is None:
                continue
            k = str(r[0]).strip().lower()
            v = "" if len(r) < 2 or r[1] is None else str(r[1]).strip()
            _put_meta(meta, k, v)
    # R6-A-003: procesar TODAS las hojas de datos con encabezados reconocibles (no solo la primera).
    # Un xlsx multi-hoja (resumen+detalle, una pestaña por tarjeta/cuenta) perdía movimientos en
    # silencio con el `break`. El locator lleva el nombre de hoja, así que los ids no colisionan.
    rows, warnings = [], []
    data_sheets = []
    for sname in wb.sheetnames:
        if sname.lower() == "meta":
            continue
        ws = wb[sname]
        it = ws.iter_rows(values_only=True)
        try:
            header = next(it)
        except StopIteration:
            continue
        cmap = {i: _canon(str(h) if h is not None else "") for i, h in enumerate(header)}
        if "date_op" not in cmap.values() or "amount" not in cmap.values():
            continue
        data_sheets.append(sname)
        for ridx, raw in enumerate(it, start=1):
            if raw is None or not any(c is not None and str(c).strip() for c in raw):
                continue
            rec = {"locator": f"{sname}!row:{ridx}", "row_idx": ridx, "table_idx": len(data_sheets) - 1,
                   "raw_text": " | ".join("" if c is None else str(c) for c in raw)}
            for i, cell in enumerate(raw):
                field = cmap.get(i)
                if field and cell is not None:
                    rec[field] = str(cell).strip()
            if "date_op" not in rec or "amount" not in rec:
                warnings.append(f"{sname}!row:{ridx} sin fecha o importe")
            rows.append(rec)
    if len(data_sheets) > 1:
        warnings.append(f"xlsx con {len(data_sheets)} hojas de datos combinadas: {', '.join(data_sheets)} "
                        f"(revisa que no haya solape resumen/detalle)")
    if account_hint and "account_id" not in meta:
        meta["account_id"] = account_hint
    return {"extraction_method": "xlsx", "extractor_version": EXTRACTOR_VERSION,
            "meta": meta, "rows": rows, "warnings": warnings}


def _put_meta(meta, k, v):
    if k == "opening_balance":
        meta["opening_raw"] = v
    elif k == "closing_balance":
        meta["closing_raw"] = v
    elif k == "declared_credits":
        meta["declared"]["credits"] = v
    elif k == "declared_debits":
        meta["declared"]["debits"] = v
    elif k == "declared_count":
        meta["declared"]["n"] = v
    elif k in META_KEYS:
        meta[k] = v
