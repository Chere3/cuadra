# -*- coding: utf-8 -*-
"""
test_r6_integridad.py — Regresión de integridad ronda 6.

- R6-A-001: un movimiento retrofechado a un MES CERRADO bloquea el commit (por date_op, no periodo).
- R6-A-002: un `type` corregido a un valor fuera de catálogo bloquea el commit.
- R6-A-003: xlsx multi-hoja procesa TODAS las hojas de datos (no pierde en silencio).
- R6-A-004: corregir categoría fija category_source='manual' (se aprende).
- R6-A-005: homóglifos (cirílico) no evaden las reglas de comercio (merchant_norm plegado).
"""
import os, sys, json, csv, shutil, sqlite3, tempfile

SRC_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def fresh():
    d = tempfile.mkdtemp(prefix="r6-")
    for sub in ("src", "config", "fixtures", "schemas"):
        shutil.copytree(os.path.join(SRC_ROOT, sub), os.path.join(d, sub))
    os.makedirs(os.path.join(d, "data"), exist_ok=True)
    return d


def load(root):
    os.environ["FINANCE_ROOT"] = root
    for n in list(sys.modules):
        if n == "finance" or n.startswith("finance."):
            del sys.modules[n]
    sys.path.insert(0, os.path.join(root, "src"))
    from finance import services as S, config as C
    C.reset_cache()
    return S


def _csv(root, name, rows, account="bbva_debito", close="0.00", credits="0.00", debits="0.00", n=1,
         pstart="2026-06-01", pend="2026-06-30"):
    p = os.path.join(root, "fixtures", name)
    with open(p, "w", encoding="utf-8") as f:
        f.write(f"# account_id: {account}\n# currency: MXN\n# period_start: {pstart}\n"
                f"# period_end: {pend}\n# opening_balance: 0.00\n# closing_balance: {close}\n"
                f"# declared_credits: {credits}\n# declared_debits: {debits}\n# declared_count: {n}\n")
        f.write("fecha,descripcion,monto,ref\n")
        for (fecha, desc, monto, ref) in rows:
            f.write(f"{fecha},{desc},{monto},{ref}\n")
    return p


def q(root, sql):
    c = sqlite3.connect(os.path.join(root, "data", "finance.sqlite"))
    c.row_factory = sqlite3.Row
    r = c.execute(sql).fetchall(); c.close(); return r


RESULTS = []
def check(name, ok, detail=""):
    RESULTS.append((name, ok, detail))
    print(("PASS " if ok else "FAIL ") + name + ((" :: " + detail) if detail else ""))


# ---------------------------------------------------------------------------
def case_closed_month_backdate():
    root = fresh(); S = load(root)
    # Junio: un movimiento, commit, cerrar el mes.
    jun = _csv(root, "jun.csv", [("2026-06-15", "NOMINA", "30000.00", "N1")],
               close="30000.00", credits="30000.00", n=1)
    S.ingest([jun], dry_run=True, run_id="jun"); S.commit("jun")
    S.close_month("2026-06")
    # Julio (mes abierto) con una línea RETROFECHADA a junio (mes cerrado).
    jul = _csv(root, "jul.csv", [("2026-07-10", "COMPRA JUL", "-1000.00", "J1"),
                                 ("2026-06-20", "COMPRA RETRO", "-5000.00", "J2")],
               close="-6000.00", debits="-6000.00", n=2, pstart="2026-07-01", pend="2026-07-31")
    S.ingest([jul], dry_run=True, run_id="jul"); r = S.commit("jul")
    jun_after = q(root, "select count(*) c, coalesce(sum(amount_minor),0) s from transactions "
                        "where substr(date_op,1,7)='2026-06'")[0]
    # junio cerrado NO debe mutar: sigue con 1 txn de 3,000,000 centavos
    ok = jun_after["c"] == 1 and jun_after["s"] == 3000000
    check("R6-A-001 mes cerrado no muta por movimiento retrofechado",
          ok, f"jun_txns={jun_after['c']} jun_sum={jun_after['s']} (esp 1/3000000)")


def case_type_out_of_catalog():
    root = fresh(); S = load(root)
    v = _csv(root, "t.csv", [("2026-06-07", "OXXO", "-250.00", "A1")],
             close="-250.00", debits="-250.00", n=1)
    S.ingest([v], dry_run=True, run_id="t")
    # forjar type basura en el staging por la vía sancionada (correct re-sella)
    p = os.path.join(root, "staging", "t", "_full.json")
    d = json.load(open(p))
    d["statements"][0]["txns"][0]["type"] = "ESTO_NO_ES_UN_TIPO"
    d["_seal"] = S._seal_staging(d)  # re-sella (simula correct sancionado)
    open(p, "w").write(json.dumps(d))
    blocked = False
    try:
        S.commit("t")
    except Exception:
        blocked = True
    ntx = q(root, "select count(*) c from transactions")[0]["c"]
    check("R6-A-002 type fuera de catálogo bloquea commit", blocked and ntx == 0,
          f"blocked={blocked} txns={ntx}")


def case_xlsx_multisheet():
    root = fresh(); S = load(root)
    import openpyxl
    wb = openpyxl.Workbook()
    meta = wb.active; meta.title = "meta"
    for k, val in [("account_id", "bbva_debito"), ("currency", "MXN"),
                   ("period_start", "2026-06-01"), ("period_end", "2026-06-30")]:
        meta.append([k, val])
    a = wb.create_sheet("cuenta_A"); a.append(["date_op", "description", "amount"])
    a.append(["2026-06-05", "MOV A1", "-100.00"]); a.append(["2026-06-06", "MOV A2", "-200.00"])
    b = wb.create_sheet("cuenta_B"); b.append(["date_op", "description", "amount"])
    b.append(["2026-06-07", "MOV B1", "-300.00"]); b.append(["2026-06-08", "MOV B2", "-400.00"])
    xp = os.path.join(root, "fixtures", "multi.xlsx"); wb.save(xp)
    from finance.extractors.xlsx_ext import extract_xlsx
    ext = extract_xlsx(xp, account_hint="bbva_debito")
    ndesc = len(ext["rows"])
    warned = any("hojas de datos" in w for w in ext["warnings"])
    check("R6-A-003 xlsx multi-hoja procesa todas + avisa", ndesc == 4 and warned,
          f"rows={ndesc} (esp 4) warned={warned}")


def case_category_source_manual():
    root = fresh(); S = load(root)
    # aplica corrección de categoría por la vía de mcp_tools y verifica category_source
    v = _csv(root, "c.csv", [("2026-06-07", "CAFE BARISTA SUR", "-80.00", "A1")],
             close="-80.00", debits="-80.00", n=1)
    S.ingest([v], dry_run=True, run_id="c")
    from finance import mcp_tools as MT
    MT.finance_apply_corrections("c", [{"transaction_id":
        json.load(open(os.path.join(root, "staging", "c", "_full.json")))["statements"][0]["txns"][0]["transaction_id"],
        "field": "category", "value": "Restaurantes/Delivery"}])
    d = json.load(open(os.path.join(root, "staging", "c", "_full.json")))
    src = d["statements"][0]["txns"][0]["category_source"]
    check("R6-A-004 corrección de categoría marca source=manual", src == "manual",
          f"category_source={src!r}")


def case_homoglyph_folding():
    root = fresh(); load(root)
    from finance import normalize as N
    latin = N.normalize_merchant("CASINO ROYALE")
    homo = N.normalize_merchant("CаSINO ROYALE")  # а cirílica
    check("R6-A-005 homóglifo se pliega a latín", latin == homo == "casino royale",
          f"latin={latin!r} homo={homo!r}")


if __name__ == "__main__":
    for fn in (case_closed_month_backdate, case_type_out_of_catalog, case_xlsx_multisheet,
               case_category_source_manual, case_homoglyph_folding):
        try:
            fn()
        except Exception as e:
            check(fn.__name__, False, "EXCEPTION " + repr(e))
    ok = all(r[1] for r in RESULTS)
    print("\n=== " + ("TODO VERDE" if ok else "HAY FALLAS") + " ===")
    sys.exit(0 if ok else 1)
